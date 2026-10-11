"""Build an editable academic manuscript from the retained Markdown source."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'Residual_Loyalty_Manuscript.md'
OUT=ROOT/'Residual_Loyalty_Manuscript.docx'
d=Document();sec=d.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.78);sec.bottom_margin=Inches(.75)
sec.left_margin=sec.right_margin=Inches(.83)
sec.header_distance=sec.footer_distance=Inches(.3)
WIDTH=6.84
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    st=d.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
    fonts=st._element.get_or_add_rPr().rFonts
    for attr in list(fonts.attrib):
        if 'Theme' in attr:del fonts.attrib[attr]
    for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),'Times New Roman')
    for node in list(st._element.xpath('.//w:pBdr')):node.getparent().remove(node)
n=d.styles['Normal'];n.font.size=Pt(11);n.paragraph_format.line_spacing=1.08;n.paragraph_format.space_after=Pt(6)
n.paragraph_format.widow_control=True
for name,size,before,after in [('Heading 1',14,14,6),('Heading 2',11.7,10,4),('Heading 3',11,8,4)]:
    st=d.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.space_before=Pt(before);st.paragraph_format.space_after=Pt(after);st.paragraph_format.keep_with_next=True
d.styles['Title'].font.size=Pt(22);d.styles['Title'].font.bold=True
d.styles['Title'].paragraph_format.space_after=Pt(5)
d.styles['Subtitle'].font.size=Pt(14);d.styles['Subtitle'].paragraph_format.space_after=Pt(13)
d.styles['Caption'].font.size=Pt(9.5);d.styles['Caption'].font.italic=False;d.styles['Caption'].font.bold=False;d.styles['Caption'].paragraph_format.line_spacing=1.02;d.styles['Caption'].paragraph_format.space_after=Pt(9)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run();r.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
d.core_properties.title='Historical Attribution Is Not Behavioral Persistence'
d.core_properties.subject='Controlled audits of learned company preferences'
d.core_properties.author='Alfaxad Eyembe'
d.core_properties.keywords='language models; alignment auditing; interpretability; remediation; Jacobian lens'

def inline(p,text):
    # Actual external hyperlinks in the editable bibliography.
    for part in re.split(r'(https?://\S+)',text):
        if part.startswith('http'):
            rel=p.part.relate_to(part,RT.HYPERLINK,is_external=True)
            link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel)
            rr=OxmlElement('w:r');pr=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Times New Roman');fonts.set(qn('w:hAnsi'),'Times New Roman');pr.append(fonts)
            size=OxmlElement('w:sz');size.set(qn('w:val'),'19');pr.append(size);rr.append(pr)
            tt=OxmlElement('w:t');tt.text=part;rr.append(tt);link.append(rr);p._p.append(link)
        else:
            for bit in re.split(r'(c_min|c_max|h_J|h_R|Δ_post|Δ_pre)',part):
                if re.fullmatch(r'(c_min|c_max|h_J|h_R|Δ_post|Δ_pre)',bit):
                    base,sub=bit.split('_');p.add_run(base).italic=True;p.add_run(sub).font.subscript=True
                else:p.add_run(bit)

def table(block):
    rows=[[c.strip() for c in line.strip().strip('|').split('|')] for line in block if not re.fullmatch(r'[\s|:\-]+',line)]
    headers=rows[0];n=len(headers)
    if headers[0]=='Population':widths=[1.38,.70,.78,3.98]
    elif headers[0]=='Intervention':widths=[1.65,.65,1.08,1.22,1.24,1.0]
    elif headers[0]=='Representation':widths=[2.79,1.35,1.35,1.35]
    elif headers[0]=='Checkpoint' and n==4:widths=[1.30,2.18,1.95,1.41]
    elif headers[0]=='Checkpoint' and n==6:widths=[.82,1.72,1.72,.86,.86,.86]
    elif headers[0]=='Pair':widths=[.42,1.13,1.03,.90,1.68,1.68]
    elif headers[0]=='Layer':widths=[.52,.92,1.08,1.22,1.55,1.55]
    elif headers[0]=='Recipe':widths=[1.74,1.02,1.02,1.02,1.02,1.02]
    elif headers[0]=='Seed':widths=[.64,2.05,2.05,2.10]
    elif headers[0]=='Date':widths=[1.05,1.30,4.49]
    else:widths=[WIDTH/n]*n
    assert abs(sum(widths)-WIDTH)<.01,(headers,widths)
    t=d.add_table(rows=0,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','bottom']:
        el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'6');el.set(qn('w:color'),'222222');borders.append(el)
    t._tbl.tblPr.append(borders)
    for ri,values in enumerate(rows):
        cells=t.add_row().cells
        for ci,(cell,text,w) in enumerate(zip(cells,values,widths)):
            cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.02;p.paragraph_format.keep_with_next=ri<len(rows)-1
            if ri==0 and text=='Independent seed groups':text='Training\nseed\ngroups'
            if ri==0 and text=='Capability loss pp':text='Capability\nloss (pp)'
            if ri==0 and text=='Bounded joint criterion':text='Bounded\njoint\ncriterion'
            inline(p,text)
            for rr in p.runs:rr.font.size=Pt(9.2);rr.bold=ri==0
            if ci and headers[0] not in ['Population','Checkpoint']:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            pr=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for side,sz in [('top',75),('bottom',75),('left',70),('right',70)]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),str(sz));e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar)
            if ri==0:
                bb=OxmlElement('w:tcBorders');e=OxmlElement('w:bottom');e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'777777');bb.append(e);pr.append(bb)
        tr=t.rows[-1]._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
        if ri==0:tr.append(OxmlElement('w:tblHeader'))
    p=d.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(2);p.add_run().font.size=Pt(2)

lines=SOURCE.read_text().splitlines();i=0;references=False;appendix=False
while i<len(lines):
    line=lines[i].strip();i+=1
    if not line:continue
    if line.startswith('# '):
        p=d.add_paragraph(line[2:],style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p=d.add_paragraph('Alfaxad Eyembe');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(3)
        p=d.add_paragraph('alfaxad@nadhari.ai');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(12)
        for rr in p.runs:rr.font.size=Pt(9.5)
        continue
    if line.startswith('## Controlled Audits'):
        continue
    if line.startswith('#'):
        h=line.lstrip('#').strip();level=1 if line.startswith('## ') else 2
        if h.startswith('Appendix '):appendix=True
        # Academic numbered hierarchy; styles remain accessible and black.
        p=d.add_paragraph(h,style='Heading '+str(level))
        if h=='Abstract':p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if h in ['References','Appendix A Additional measurement details']:p.paragraph_format.page_break_before=True
        references=h=='References'
        continue
    if line.startswith('|'):
        block=[line]
        while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i].strip());i+=1
        table(block);continue
    if line.startswith('!['):
        m=re.match(r'!\[([^\]]*)\]\(([^)]+)\)',line);p=d.add_paragraph();p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3)
        pic=p.add_run().add_picture(str(ROOT/m[2]),width=Inches(WIDTH));pic._inline.docPr.set('descr',m[1]);continue
    if re.match(r'^(Figure|Table) [A-Z]?\d+\.',line):
        p=d.add_paragraph(style='Caption');p.paragraph_format.keep_with_next=line.startswith('Table ');inline(p,line);continue
    if re.search(r'\([123]\)$',line):
        p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(9)
        inline(p,re.sub(r' {2,}','   ',line))
        for rr in p.runs:rr.font.name='Cambria Math';rr.font.size=Pt(11)
        continue
    p=d.add_paragraph();inline(p,line)
    if references:
        p.paragraph_format.space_after=Pt(9);p.paragraph_format.line_spacing=1.0
        p.paragraph_format.left_indent=Inches(.25);p.paragraph_format.first_line_indent=Inches(-.25)
        for rr in p.runs:rr.font.size=Pt(9.5)
    else:
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        if appendix:
            p.paragraph_format.line_spacing=1.04;p.paragraph_format.space_after=Pt(5)
            for rr in p.runs:rr.font.size=Pt(10.5)
d.save(OUT)
print(OUT)
